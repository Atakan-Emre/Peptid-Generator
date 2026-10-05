#!/usr/bin/env python3
"""Hakem cevap belgelerini ve makale revizyon notunu uretir.

    python scripts/make_review_documents.py [--out ../HakemCevaplari]

Uc dosya yazar:
    Hakem_Cevap_1_Reviewer_1.docx   Hakem 1'in 4 maddesine nokta nokta cevap
    Hakem_Cevap_2_Reviewer_2.docx   Hakem 2'nin 6 maddesine nokta nokta cevap
    Makale_Revizyon_Notlari.pdf     Makalede sayfa sayfa ne degisecegi

Bicim, HakemCevaplari klasorundeki MDPI ornek belgesinden alinmistir:
Times New Roman, A4, "Comment N:" / "Response:" kalip.

Cevaplardaki her sayi results altindaki artefaktlardan gelir; dosya
icindeki NUMBERS sozlugu tek kaynaktir, boylece belge ile depo ayrisamaz.
"""
import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

ROOT = Path(__file__).resolve().parents[1]

PAPER_TITLE = ("AI-driven design of plastic-binding peptides toward analytical "
               "recognition of multiple microplastic types")
JOURNAL = "[Journal name]"
MS_ID = "[Manuscript ID]"


# --------------------------------------------------------------------------
# Sayilar tek yerden: artefaktlardan okunur, elle yazilmaz.
# --------------------------------------------------------------------------
def load_numbers(out_dir: Path) -> dict:
    j = lambda p: json.loads((out_dir / p).read_text(encoding="utf-8"))
    clu = j("summary_clustered.json")
    rnd = j("summary_random.json")
    st = j("architecture_stats_clustered.json")
    jain = j("jain_baseline/jain_karsilastirma.json")
    sel = j("analysis/selectivity.json")
    nov = j("analysis/novelty.json")
    iv = j("analysis/independent_validation.json")
    gen = j("analysis/generator_comparison.json")

    def mean_r2(data, arch):
        v = [r["test_r2_mean"] for r in data if r["architecture"] == arch]
        return sum(v) / len(v)

    da = [jain[p]["OURS_A_ham_rastgele"]["test_r2_mean"]
          - jain[p]["A_ham_rastgele"]["test_r2_mean"] for p in jain]
    dc = [jain[p]["OURS_C_temiz_kumelenmis"]["test_r2_mean"]
          - jain[p]["C_temiz_kumelenmis"]["test_r2_mean"] for p in jain]

    return {
        "cnn_clustered": mean_r2(clu, "cnn"),
        "cnn_random": mean_r2(rnd, "cnn"),
        "encdec_clustered": mean_r2(clu, "encdec"),
        "n_comparisons": st["summary"]["n_comparisons"],
        "n_significant": st["summary"]["n_significant"],
        "jain_gain_A": sum(da) / len(da),
        "jain_gain_C": sum(dc) / len(dc),
        "jain_drop_pct": 100 * (1 - (sum(dc) / len(dc)) / (sum(da) / len(da))),
        "sel_target_best": sel["summary"]["n_target_best"],
        "sel_n": sel["summary"]["n_targets"],
        "sel_min": min(sel["matrix"][t]["selectivity_index"] for t in sel["plastics"]),
        "sel_max": max(sel["matrix"][t]["selectivity_index"] for t in sel["plastics"]),
        "nov_min": min(r["nn_identity_mean_pct"] for r in nov),
        "nov_max": max(r["nn_identity_mean_pct"] for r in nov),
        "nov_above90": max(r["pct_above_90_identity"] for r in nov),
        "iv_beats": sum(1 for r in iv if r["beats_best_training_sequence"]),
        "iv_confirms": sum(1 for r in iv if r["validator_confirms_gain"]),
        "iv_n": len(iv),
        "validator_arch": iv[0]["validator_model"]["architecture"],
        "selected_generator": gen["selection"]["selected"],
        "clu": {r["plastic"]: r for r in clu if r["architecture"] == "cnn"},
        "clu_all": clu,
        "iv": {r["plastic"]: r for r in iv},
        "nov": {r["plastic"]: r for r in nov},
        "sel": sel,
        "jain": jain,
    }


# --------------------------------------------------------------------------
# docx uretimi
# --------------------------------------------------------------------------
def build_docx(path: Path, reviewer_no: int, intro: str, items: list) -> None:
    import docx
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.5)

    def para(text="", bold=False, italic=False, space_after=6, align=None):
        p = d.add_paragraph()
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        p.paragraph_format.space_after = Pt(space_after)
        if align:
            p.alignment = align
        return p

    h = d.add_heading("Authors' Response to the Review Comments", level=1)
    h.runs[0].font.name = "Times New Roman"

    para(f"Journal\t\t: {JOURNAL}", bold=True, space_after=0)
    para(f"Manuscript ID\t: {MS_ID}", bold=True, space_after=0)
    para(f"Title of Paper\t: {PAPER_TITLE}", bold=True, space_after=12)

    para(intro, space_after=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    hh = d.add_heading(f"Responses to Comments of Reviewer #{reviewer_no}", level=1)
    hh.runs[0].font.name = "Times New Roman"

    for n, item in enumerate(items, 1):
        para(f"Comment {n}: {item['comment']}", bold=True, space_after=6,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        para("Response:", bold=True, space_after=6)
        for block in item["response"]:
            if isinstance(block, dict) and block.get("table"):
                _add_table(d, block["table"], block.get("caption"))
            else:
                para(block, space_after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        para("", space_after=6)

    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)


def _add_table(d, rows, caption=None):
    from docx.shared import Pt
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for jx, cell in enumerate(row):
            c = t.cell(i, jx)
            c.text = ""
            p = c.paragraphs[0]
            r = p.add_run(str(cell))
            r.font.size = Pt(9)
            r.font.name = "Times New Roman"
            r.bold = (i == 0)
            p.paragraph_format.space_after = Pt(0)
    if caption:
        p = d.add_paragraph()
        r = p.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        p.paragraph_format.space_after = Pt(10)
    else:
        d.add_paragraph().paragraph_format.space_after = Pt(10)


# --------------------------------------------------------------------------
# PDF uretimi (makale revizyon notu)
# --------------------------------------------------------------------------
def build_pdf(path: Path, sections: list, summary: list) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, KeepTogether)

    ss = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=ss["Title"], fontName="Times-Bold",
                           fontSize=16, spaceAfter=4)
    sub = ParagraphStyle("s", parent=ss["Normal"], fontName="Times-Italic",
                         fontSize=10, textColor=colors.HexColor("#566573"),
                         alignment=1, spaceAfter=14)
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Times-Roman",
                          fontSize=10, leading=14, spaceAfter=8, alignment=4)
    hdr = ParagraphStyle("h", parent=ss["Heading2"], fontName="Times-Bold",
                         fontSize=12, textColor=colors.HexColor("#1b4f72"),
                         spaceBefore=14, spaceAfter=6)
    cell = ParagraphStyle("c", parent=ss["Normal"], fontName="Times-Roman",
                          fontSize=8.5, leading=11)
    cellb = ParagraphStyle("cb", parent=cell, fontName="Times-Bold",
                           textColor=colors.white)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=A4,
                            leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm,
                            title="Manuscript revision notes")
    flow = [Paragraph("Manuscript Revision Notes", title),
            Paragraph(PAPER_TITLE, sub)]

    for block in summary:
        flow.append(Paragraph(block, body))

    widths = [1.3 * cm, 3.1 * cm, 5.4 * cm, 5.4 * cm, 2.1 * cm]
    for sec in sections:
        flow.append(Paragraph(sec["title"], hdr))
        if sec.get("note"):
            flow.append(Paragraph(sec["note"], body))
        data = [[Paragraph(h, cellb) for h in
                 ("Page", "Location", "What it says now", "What it must become",
                  "Evidence")]]
        for r in sec["rows"]:
            data.append([Paragraph(str(x), cell) for x in r])
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1b4f72")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#aab7b8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#f4f6f7")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        flow.append(t)
        flow.append(Spacer(1, 4))

    doc.build(flow)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT.parent / "HakemCevapları"))
    ap.add_argument("--results", default=str(ROOT / "results"))
    a = ap.parse_args()

    out = Path(a.out)
    n = load_numbers(Path(a.results))

    from review_content import reviewer1, reviewer2, intro1, intro2, pdf_sections, pdf_summary

    build_docx(out / "Hakem_Cevap_1_Reviewer_1.docx", 1, intro1(n), reviewer1(n))
    print(f"  {out / 'Hakem_Cevap_1_Reviewer_1.docx'}")
    build_docx(out / "Hakem_Cevap_2_Reviewer_2.docx", 2, intro2(n), reviewer2(n))
    print(f"  {out / 'Hakem_Cevap_2_Reviewer_2.docx'}")
    build_pdf(out / "Makale_Revizyon_Notlari.pdf", pdf_sections(n), pdf_summary(n))
    print(f"  {out / 'Makale_Revizyon_Notlari.pdf'}")
    return 0


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    raise SystemExit(main())
